# -*- coding: utf-8 -*-
"""Genera el Excel de deudas de UN productor a partir de deudas_full.json (openpyxl).

Uso:
  python build_xlsx.py
  python build_xlsx.py --in salida/7224/deudas_full.json --out salida/7224/deudas.xlsx
"""
import argparse, json, re, datetime, os
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule, FormulaRule

ap = argparse.ArgumentParser()
ap.add_argument('--in', dest='inp', default='deudas_full.json')
ap.add_argument('--out', dest='out', default='deudas.xlsx')
ap.add_argument('--meta', default=None)
args = ap.parse_args()

with open(args.inp, encoding='utf-8') as f:
    data = json.load(f)

META = {}
if args.meta and os.path.exists(args.meta):
    with open(args.meta, encoding='utf-8') as f:
        META = json.load(f)

HOY = datetime.date.today()   # fecha de generación del listado
AZUL = '1F3864'
AZUL_CLARO = 'D9E2F3'
GRIS = 'F2F2F2'
VERDE = 'E2EFDA'
ROJO = 'FFC7CE'
AMBAR = 'FFEB9C'

thin = Side(style='thin', color='BFBFBF')
BORDER = Border(bottom=thin)

SIMBOLOS = {
    'R': 'Sin cobertura financiera por falta de pago',
    'V': 'Deuda vencida dentro del plazo adicional',
    'A': 'Con deuda vencida',
    '+': 'Más de una cuota en el día',
    '.': 'Pagos en trámite',
    '1': 'Primera cuota de la póliza',
    'G': 'Cuota garantizada',
}

def situacion(marca):
    partes, vistos = [], set()
    for ch in (marca or ''):
        txt = SIMBOLOS.get(ch)
        if txt and txt not in vistos:
            partes.append(txt); vistos.add(txt)
    return ' + '.join(partes) if partes else ''

def ar(n):
    return f'{n:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')

def tel_e164(t):
    raw = re.sub(r'\D', '', t.get('raw', ''))
    if not raw:
        return '', ''
    d = raw[1:] if raw.startswith('0') else raw
    if t.get('celular'):
        for L in (4, 3, 2):
            if len(d) > L + 2 and d[L:L + 2] == '15':
                if L == 2 and d[:2] != '11':
                    continue
                resto = d[L + 2:]
                if 6 <= len(resto) <= 8:
                    return '+549' + d[:L] + resto, 'movil'
    return '+54' + d, 'fijo'

def fmt_tel(t):
    raw = re.sub(r'\D', '', t.get('raw', ''))
    return raw

def fecha_iso(q):
    try:
        return datetime.date.fromisoformat(q['vencimientoISO'])
    except Exception:
        return None

# ---------- preparar filas ----------
filas = []
for q in data['cuotas']:
    f = fecha_iso(q)
    tels = q.get('telefonos') or []
    t = next((x for x in tels if x.get('celular')), tels[0] if tels else None)
    wa, tipo_tel = tel_e164(t) if t else ('', '')
    mails = [m['mail'] for m in (q.get('mails') or [])]
    mails = list(dict.fromkeys(mails))
    dias = (HOY - f).days if f else None
    filas.append({
        'asegurado': q['asegurado'],
        'matricula': q['matricula'],
        'poliza': q.get('poliza', ''),
        'cuota': f"{q['cuotaNumero']}/{q['cuotaTotal']}" if q.get('cuotaNumero') else '',
        'vencimiento': f,
        'dias': dias,
        'saldo': float(q.get('saldo') or 0),
        'estado': q.get('estado', ''),
        'marca': q.get('marca', ''),
        'situacion': situacion(q.get('marca', '')),
        'telefono': fmt_tel(t) if t else '',
        'whatsapp': wa,
        'tipo_tel': tipo_tel,
        'email': mails[0] if mails else '',
        'otros_email': '; '.join(mails[1:]),
        'nroOp': q.get('nroOperacion', ''),
        'id': q.get('cajaId', ''),
    })

filas.sort(key=lambda r: (r['vencimiento'] or datetime.date.max, r['asegurado']))

wb = Workbook()

# =================== HOJA 1: DEUDAS ===================
ws = wb.active
ws.title = 'Deudas'
ws.sheet_properties.tabColor = AZUL

COLS = [
    ('Asegurado', 34, 'asegurado'),
    ('Matrícula', 11, 'matricula'),
    ('Póliza', 15, 'poliza'),
    ('Cuota', 8, 'cuota'),
    ('Vencimiento', 12, 'vencimiento'),
    ('Días', 8, 'dias'),
    ('Saldo', 14, 'saldo'),
    ('Estado', 10, 'estado'),
    ('Situación', 34, 'situacion'),
    ('Teléfono', 15, 'telefono'),
    ('WhatsApp', 16, 'whatsapp'),
    ('Email', 30, 'email'),
    ('Otros emails', 28, 'otros_email'),
    ('Link WhatsApp', 24, 'wa_link'),
    ('ID cuota SELF', 26, 'id'),
]

ws.append([c[0] for c in COLS])
for i, (_, w, _) in enumerate(COLS, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

for r in filas:
    r['wa_link'] = f"https://wa.me/{r['whatsapp'].replace('+', '')}" if r['tipo_tel'] == 'movil' else ''
    ws.append([r[c[2]] for c in COLS])

n = len(filas)
last_row = n + 1

# estilos de encabezado
for c in ws[1]:
    c.fill = PatternFill('solid', fgColor=AZUL)
    c.font = Font(bold=True, color='FFFFFF', size=10)
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws.row_dimensions[1].height = 30

# formatos + zebra
for idx in range(2, last_row + 1):
    rec = filas[idx - 2]
    for j, (_, _, key) in enumerate(COLS, start=1):
        cell = ws.cell(row=idx, column=j)
        cell.font = Font(size=10)
        cell.border = BORDER
        cell.alignment = Alignment(vertical='center', wrap_text=key in ('situacion', 'asegurado', 'otros_email'))
        if key == 'saldo':
            cell.number_format = '"$" #,##0.00'
            cell.alignment = Alignment(horizontal='right', vertical='center')
        elif key == 'vencimiento' and cell.value:
            cell.number_format = 'dd/mm/yyyy'
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif key in ('dias', 'cuota', 'estado', 'marca', 'tipo_tel', 'nroOp'):
            cell.alignment = Alignment(horizontal='center', vertical='center')
        if key == 'wa_link' and cell.value:
            cell.hyperlink = cell.value
            cell.value = 'Abrir chat'
            cell.font = Font(size=10, color='0563C1', underline='single')
    if rec['estado'] == 'ANULADA':
        for j in range(1, len(COLS) + 1):
            ws.cell(row=idx, column=j).fill = PatternFill('solid', fgColor=GRIS)
            ws.cell(row=idx, column=j).font = Font(size=10, color='808080', italic=True)
    elif idx % 2 == 0:
        for j in range(1, len(COLS) + 1):
            ws.cell(row=idx, column=j).fill = PatternFill('solid', fgColor='F7F9FC')

# colores en Días (vencidas)
ws.conditional_formatting.add(f'F2:F{last_row}',
    CellIsRule(operator='greaterThan', formula=['30'], fill=PatternFill('solid', fgColor=ROJO), font=Font(color='9C0006', bold=True)))
ws.conditional_formatting.add(f'F2:F{last_row}',
    CellIsRule(operator='between', formula=['1', '30'], fill=PatternFill('solid', fgColor=AMBAR)))
ws.conditional_formatting.add(f'F2:F{last_row}',
    CellIsRule(operator='lessThan', formula=['1'], fill=PatternFill('solid', fgColor=VERDE)))

# fila de total
tot_row = last_row + 1
ws.cell(row=tot_row, column=1, value=f'TOTAL — {n} cuotas impagas').font = Font(bold=True, size=11)
ws.merge_cells(start_row=tot_row, start_column=1, end_row=tot_row, end_column=6)
ws.cell(row=tot_row, column=1).alignment = Alignment(horizontal='right')
tc = ws.cell(row=tot_row, column=7, value=sum(r['saldo'] for r in filas))
tc.number_format = '"$" #,##0.00'
tc.font = Font(bold=True, size=11, color=AZUL)
tc.alignment = Alignment(horizontal='right')
for j in range(1, len(COLS) + 1):
    ws.cell(row=tot_row, column=j).fill = PatternFill('solid', fgColor=AZUL_CLARO)
    ws.cell(row=tot_row, column=j).border = Border(top=Side(style='double', color=AZUL))

ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:{get_column_letter(len(COLS))}{last_row}'
ws.sheet_view.zoomScale = 90
ws.page_setup.orientation = 'landscape'
ws.page_setup.fitToWidth = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = '1:1'

# =================== HOJA 2: POR CLIENTE ===================
ws2 = wb.create_sheet('Por cliente')
ws2.sheet_properties.tabColor = '2E75B6'
por_cli = defaultdict(lambda: {'polizas': set(), 'cuotas': 0, 'total': 0.0, 'venc': None,
                               'tel': '', 'wa': '', 'email': '', 'mat': '', 'anuladas': 0})
for r in filas:
    c = por_cli[r['asegurado']]
    c['mat'] = r['matricula']
    if r['poliza']:
        c['polizas'].add(r['poliza'])
    c['cuotas'] += 1
    c['total'] += r['saldo']
    if r['estado'] == 'ANULADA':
        c['anuladas'] += 1
    if r['vencimiento'] and (c['venc'] is None or r['vencimiento'] < c['venc']):
        c['venc'] = r['vencimiento']
    c['tel'] = c['tel'] or r['telefono']
    c['wa'] = c['wa'] or r['whatsapp']
    c['email'] = c['email'] or r['email']

cols2 = [('Asegurado', 34,), ('Matrícula', 11), ('Pólizas', 26), ('Cuotas', 9),
         ('Total adeudado', 15), ('Venc. más antiguo', 14), ('Días', 8),
         ('Teléfono', 15), ('WhatsApp', 16), ('Email', 32), ('Link WhatsApp', 18)]
ws2.append([c[0] for c in cols2])
for i, c in enumerate(cols2, start=1):
    ws2.column_dimensions[get_column_letter(i)].width = c[1]

orden = sorted(por_cli.items(), key=lambda kv: -kv[1]['total'])
for nombre, c in orden:
    dias = (HOY - c['venc']).days if c['venc'] else None
    ws2.append([
        nombre, f"M-{c['mat']}", ', '.join(sorted(c['polizas'])), c['cuotas'], c['total'],
        c['venc'], dias, c['tel'], c['wa'], c['email'],
        f"https://wa.me/{c['wa'].replace('+','')}" if c['wa'].startswith('+549') else '',
    ])

last2 = len(orden) + 1
for cl in ws2[1]:
    cl.fill = PatternFill('solid', fgColor=AZUL)
    cl.font = Font(bold=True, color='FFFFFF', size=10)
    cl.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws2.row_dimensions[1].height = 30
for i in range(2, last2 + 1):
    for j in range(1, len(cols2) + 1):
        cell = ws2.cell(row=i, column=j)
        cell.font = Font(size=10)
        cell.border = BORDER
        cell.alignment = Alignment(vertical='center', wrap_text=j in (1, 3, 10))
        if j == 5:
            cell.number_format = '"$" #,##0.00'
            cell.alignment = Alignment(horizontal='right', vertical='center')
        elif j == 6:
            cell.number_format = 'dd/mm/yyyy'
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif j in (2, 4, 7, 8, 9):
            cell.alignment = Alignment(horizontal='center', vertical='center')
    if i % 2 == 0:
        for j in range(1, len(cols2) + 1):
            ws2.cell(row=i, column=j).fill = PatternFill('solid', fgColor='F7F9FC')
    wa_cell = ws2.cell(row=i, column=11)
    if wa_cell.value:
        wa_cell.hyperlink = wa_cell.value
        wa_cell.value = 'Abrir chat'
        wa_cell.font = Font(size=10, color='0563C1', underline='single')

ws2.conditional_formatting.add(f'G2:G{last2}',
    CellIsRule(operator='greaterThan', formula=['30'], fill=PatternFill('solid', fgColor=ROJO), font=Font(color='9C0006', bold=True)))
ws2.conditional_formatting.add(f'G2:G{last2}',
    CellIsRule(operator='lessThan', formula=['1'], fill=PatternFill('solid', fgColor=VERDE)))

tot2 = last2 + 1
ws2.cell(row=tot2, column=1, value=f'TOTAL — {len(orden)} clientes').font = Font(bold=True, size=11)
ws2.merge_cells(start_row=tot2, start_column=1, end_row=tot2, end_column=4)
ws2.cell(row=tot2, column=1).alignment = Alignment(horizontal='right')
t2 = ws2.cell(row=tot2, column=5, value=sum(c['total'] for _, c in orden))
t2.number_format = '"$" #,##0.00'
t2.font = Font(bold=True, size=11, color=AZUL)
t2.alignment = Alignment(horizontal='right')
for j in range(1, len(cols2) + 1):
    ws2.cell(row=tot2, column=j).fill = PatternFill('solid', fgColor=AZUL_CLARO)
    ws2.cell(row=tot2, column=j).border = Border(top=Side(style='double', color=AZUL))
ws2.freeze_panes = 'A2'
ws2.auto_filter.ref = f'A1:K{last2}'
ws2.sheet_view.zoomScale = 90
ws2.page_setup.orientation = 'landscape'
ws2.sheet_properties.pageSetUpPr.fitToPage = True
ws2.page_setup.fitToWidth = 1

# =================== HOJA 3: RESUMEN ===================
ws3 = wb.create_sheet('Resumen')
ws3.sheet_properties.tabColor = 'A9D08E'
ws3.column_dimensions['A'].width = 34
ws3.column_dimensions['B'].width = 46

def titulo(fila, texto):
    c = ws3.cell(row=fila, column=1, value=texto)
    c.font = Font(bold=True, size=13, color=AZUL)
    ws3.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=2)

def par(fila, k, v, fmt=None, bold=False):
    a = ws3.cell(row=fila, column=1, value=k)
    b = ws3.cell(row=fila, column=2, value=v)
    a.font = Font(size=10, color='595959')
    b.font = Font(size=10, bold=bold)
    if fmt:
        b.number_format = fmt
    a.border = BORDER
    b.border = BORDER

total = sum(r['saldo'] for r in filas)
total_vig = sum(r['saldo'] for r in filas if r['estado'] != 'ANULADA')
vencidas = [r for r in filas if r['dias'] is not None and r['dias'] > 0]
monto_vencido = sum(r['saldo'] for r in vencidas)

titulo(1, 'Cuotas impagas — SELF Federación Patronal')
ws3.cell(row=2, column=1, value='Generado automáticamente (extracción HTTP)').font = Font(size=9, italic=True, color='808080')

def fecha_legible(x):
    if not x:
        return ''
    p = str(x).split('/')
    if len(p) == 3:
        d, m, y = p
        if len(y) == 2:
            y = '20' + y
        return f'{d}/{m}/{y}'
    return x

PRODUCTOR = str(META.get('user') or data.get('productor') or '')
NOMBRE = META.get('nombre') or data.get('nombre') or ''
RANGO = META.get('rango') or data.get('rango') or {}

f = 4
par(f, 'Productor', f'{PRODUCTOR} — {NOMBRE}'.strip(' —')); f += 1
par(f, 'Rango del listado', f"{fecha_legible(RANGO.get('d'))} a {fecha_legible(RANGO.get('h'))}".strip(' a')); f += 1
par(f, 'Fecha de generación', datetime.datetime.now().strftime('%d/%m/%Y %H:%M')); f += 1
f += 1
par(f, 'Clientes con deuda', len(orden), bold=True); f += 1
par(f, 'Cuotas impagas', len(filas), bold=True); f += 1
par(f, 'Monto total', total, '"$" #,##0.00', bold=True); f += 1
par(f, 'Monto total (sin ANULADAS)', total_vig, '"$" #,##0.00'); f += 1
par(f, 'Cuotas ya vencidas', f'{len(vencidas)} — $ {ar(monto_vencido)}'); f += 1
par(f, 'Cuotas a vencer', len(filas) - len(vencidas)); f += 2

titulo(f, 'Composición por situación'); f += 1
por_sit = defaultdict(lambda: [0, 0.0])
for r in filas:
    por_sit[r['situacion'] or '(sin dato)'][0] += 1
    por_sit[r['situacion'] or '(sin dato)'][1] += r['saldo']
for k, (cnt, monto) in sorted(por_sit.items(), key=lambda kv: -kv[1][1]):
    par(f, k, f'{cnt} cuota(s) — $ {ar(monto)}'); f += 1
f += 1

titulo(f, 'Composición por estado'); f += 1
por_est = defaultdict(lambda: [0, 0.0])
for r in filas:
    por_est[r['estado'] or '(sin dato)'][0] += 1
    por_est[r['estado'] or '(sin dato)'][1] += r['saldo']
for k, (cnt, monto) in sorted(por_est.items(), key=lambda kv: -kv[1][1]):
    par(f, k, f'{cnt} cuota(s) — $ {ar(monto)}'); f += 1
f += 1

titulo(f, 'Referencias'); f += 1
for k, v in SIMBOLOS.items():
    par(f, f'Marca "{k}"', v); f += 1
par(f, 'Días', 'Positivo = días vencida · Negativo = días para vencer'); f += 1
par(f, 'WhatsApp', 'Formato internacional +54 9 (móviles). Verificar antes de enviar.')

# =================== guardar ===================
out = args.out
os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
wb.save(out)
print(f'OK [{PRODUCTOR}] -> {out} | {len(filas)} cuotas | {len(orden)} clientes | total $ {total:,.2f}')
