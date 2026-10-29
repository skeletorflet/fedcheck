# -*- coding: utf-8 -*-
"""Consolidado multi-cuenta: deudas_todas.xlsx

Junta los deudas_full.json de cada productor (salida/<user>/) en un solo libro:
  - Resumen      : totales por productor y globales
  - Deudas       : todas las cuotas, con columna Productor
  - Por cliente  : total adeudado por cliente, con Productor

Uso:
  python build_all.py
  python build_all.py --out deudas_todas.xlsx salida/7224/deudas_full.json ...
"""
import argparse, json, os, re, datetime, glob
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

AZUL = '1F3864'
AZUL_CLARO = 'D9E2F3'
GRIS = 'F2F2F2'
VERDE = 'E2EFDA'
ROJO = 'FFC7CE'
AMBAR = 'FFEB9C'
BORDE = Border(bottom=Side(style='thin', color='BFBFBF'))
HOY = datetime.date.today()

SIMBOLOS = {
    'R': 'Sin cobertura financiera por falta de pago',
    'V': 'Deuda vencida dentro del plazo adicional',
    'A': 'Con deuda vencida',
    '+': 'Más de una cuota en el día',
    '.': 'Pagos en trámite',
    '1': 'Primera cuota de la póliza',
    'G': 'Cuota garantizada',
}

ap = argparse.ArgumentParser()
ap.add_argument('--out', default='deudas_todas.xlsx')
ap.add_argument('files', nargs='*')
args = ap.parse_args()

files = args.files or sorted(glob.glob('salida/*/deudas_full.json'))
if not files:
    raise SystemExit('no hay deudas_full.json para consolidar')

def situacion(marca):
    partes, vistos = [], set()
    for ch in (marca or ''):
        txt = SIMBOLOS.get(ch)
        if txt and txt not in vistos:
            partes.append(txt); vistos.add(txt)
    return ' + '.join(partes)

def ar(n):
    return f'{n:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')

def tel_e164(t):
    raw = re.sub(r'\D', '', t.get('raw', ''))
    if not raw:
        return '', ''
    d = raw[1:] if raw.startswith('0') else raw
    if t.get('celular'):
        for L in (4, 3, 2):
            if len(d) > L + 2 and d[L:L + 2] == '15':
                if L == 2 and d[:2] != '11':
                    continue
                resto = d[L + 2:]
                if 6 <= len(resto) <= 8:
                    return '+549' + d[:L] + resto, 'movil'
    return '+54' + d, 'fijo'

def fecha_iso(q):
    try:
        return datetime.date.fromisoformat(q['vencimientoISO'])
    except Exception:
        return None

# ---------- cargar ----------
filas = []
productores = {}
for path in files:
    with open(path, encoding='utf-8') as f:
        data = json.load(f)
    user = str(data.get('productor') or os.path.basename(os.path.dirname(path)))
    nombre = data.get('nombre') or ''
    productores[user] = {'nombre': nombre, 'clientes': data.get('clientes', 0), 'cuotas': 0,
                         'total': 0.0, 'vencidas': 0, 'monto_vencido': 0.0, 'rango': data.get('rango') or {}}
    for q in data['cuotas']:
        fch = fecha_iso(q)
        tels = q.get('telefonos') or []
        t = next((x for x in tels if x.get('celular')), tels[0] if tels else None)
        wa, tipo = tel_e164(t) if t else ('', '')
        mails = list(dict.fromkeys(m['mail'] for m in (q.get('mails') or [])))
        dias = (HOY - fch).days if fch else None
        saldo = float(q.get('saldo') or 0)
        p = productores[user]
        p['cuotas'] += 1
        p['total'] += saldo
        if dias is not None and dias > 0:
            p['vencidas'] += 1
            p['monto_vencido'] += saldo
        filas.append({
            'productor': user, 'asegurado': q['asegurado'], 'matricula': q['matricula'],
            'poliza': q.get('poliza', ''), 'cuota': f"{q['cuotaNumero']}/{q['cuotaTotal']}" if q.get('cuotaNumero') else '',
            'vencimiento': fch, 'dias': dias, 'saldo': saldo, 'estado': q.get('estado', ''),
            'situacion': situacion(q.get('marca', '')),
            'telefono': re.sub(r'\D', '', t.get('raw', '')) if t else '',
            'whatsapp': wa, 'tipo': tipo,
            'email': mails[0] if mails else '', 'otros': '; '.join(mails[1:]),
            'wa_link': f"https://wa.me/{wa.replace('+', '')}" if tipo == 'movil' else '',
            'id': q.get('cajaId', ''),
        })

filas.sort(key=lambda r: (r['vencimiento'] or datetime.date.max, r['productor'], r['asegurado']))
total_global = sum(r['saldo'] for r in filas)

wb = Workbook()

# =================== RESUMEN ===================
ws = wb.active
ws.title = 'Resumen'
ws.sheet_properties.tabColor = 'A9D08E'
anchos = {'A': 12, 'B': 34, 'C': 11, 'D': 10, 'E': 16, 'F': 12, 'G': 16, 'H': 10}
for k, v in anchos.items():
    ws.column_dimensions[k].width = v

ws['A1'] = 'Cuotas impagas — consolidado multi-producto'
ws['A1'].font = Font(bold=True, size=14, color=AZUL)
ws.merge_cells('A1:H1')
ws['A2'] = 'SELF Federación Patronal · generado automáticamente (extracción HTTP)'
ws['A2'].font = Font(size=9, italic=True, color='808080')
ws.merge_cells('A2:H2')

encabezados = ['Productor', 'Nombre', 'Clientes', 'Cuotas', 'Total adeudado', 'Vencidas', 'A vencer', '% del total']
fila = 4
for j, h in enumerate(encabezados, start=1):
    c = ws.cell(row=fila, column=j, value=h)
    c.fill = PatternFill('solid', fgColor=AZUL)
    c.font = Font(bold=True, color='FFFFFF', size=10)
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws.row_dimensions[fila].height = 28

orden = sorted(productores.items(), key=lambda kv: -kv[1]['total'])
r = fila + 1
for user, p in orden:
    a_vencer = p['cuotas'] - p['vencidas']
    vals = [user, p['nombre'], p['clientes'], p['cuotas'], p['total'], p['vencidas'], a_vencer,
            (p['total'] / total_global) if total_global else 0]
    for j, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=j, value=v)
        c.font = Font(size=10)
        c.border = BORDE
        if j == 5:
            c.number_format = '"$" #,##0.00'
            c.alignment = Alignment(horizontal='right')
        elif j == 8:
            c.number_format = '0.0%'
            c.alignment = Alignment(horizontal='center')
        elif j in (3, 4, 6, 7):
            c.alignment = Alignment(horizontal='center')
    if r % 2 == 0:
        for j in range(1, 9):
            ws.cell(row=r, column=j).fill = PatternFill('solid', fgColor='F7F9FC')
    r += 1

# total
tot_v = sum(p['vencidas'] for _, p in orden)
tot_c = sum(p['cuotas'] for _, p in orden)
vals = ['TOTAL', f'{len(orden)} productores', sum(p['clientes'] for _, p in orden), tot_c, total_global, tot_v, tot_c - tot_v, 1]
for j, v in enumerate(vals, start=1):
    c = ws.cell(row=r, column=j, value=v)
    c.fill = PatternFill('solid', fgColor=AZUL_CLARO)
    c.font = Font(bold=True, size=11, color=AZUL)
    c.border = Border(top=Side(style='double', color=AZUL))
    if j == 5:
        c.number_format = '"$" #,##0.00'
        c.alignment = Alignment(horizontal='right')
    elif j == 8:
        c.number_format = '0.0%'
        c.alignment = Alignment(horizontal='center')

# detalle por productor (rango + vencido)
r += 2
ws.cell(row=r, column=1, value='Detalle').font = Font(bold=True, size=12, color=AZUL)
r += 1
for user, p in orden:
    rg = p.get('rango') or {}
    txt = f"{user} {p['nombre']} — rango {rg.get('d','')} a {rg.get('h','')} — vencido $ {ar(p['monto_vencido'])} de $ {ar(p['total'])}"
    c = ws.cell(row=r, column=1, value=txt)
    c.font = Font(size=10, color='595959')
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    r += 1

r += 1
ws.cell(row=r, column=1, value='Referencias').font = Font(bold=True, size=12, color=AZUL)
r += 1
for k, v in SIMBOLOS.items():
    ws.cell(row=r, column=1, value=f'Marca "{k}"').font = Font(size=9, color='595959')
    c = ws.cell(row=r, column=2, value=v)
    c.font = Font(size=9)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    r += 1

# =================== DEUDAS ===================
ws2 = wb.create_sheet('Deudas')
ws2.sheet_properties.tabColor = AZUL
COLS = [
    ('Productor', 11, 'productor'),
    ('Asegurado', 32, 'asegurado'),
    ('Matrícula', 11, 'matricula'),
    ('Póliza', 15, 'poliza'),
    ('Cuota', 8, 'cuota'),
    ('Vencimiento', 12, 'vencimiento'),
    ('Días', 8, 'dias'),
    ('Saldo', 14, 'saldo'),
    ('Estado', 10, 'estado'),
    ('Situación', 32, 'situacion'),
    ('Teléfono', 15, 'telefono'),
    ('WhatsApp', 16, 'whatsapp'),
    ('Email', 30, 'email'),
    ('Link WhatsApp', 16, 'wa_link'),
    ('ID cuota SELF', 26, 'id'),
]
ws2.append([c[0] for c in COLS])
for i, (_, w, _) in enumerate(COLS, start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w

for rec in filas:
    ws2.append([rec[c[2]] for c in COLS])

last = len(filas) + 1
for c in ws2[1]:
    c.fill = PatternFill('solid', fgColor=AZUL)
    c.font = Font(bold=True, color='FFFFFF', size=10)
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws2.row_dimensions[1].height = 30

for i in range(2, last + 1):
    rec = filas[i - 2]
    for j, (_, _, key) in enumerate(COLS, start=1):
        cell = ws2.cell(row=i, column=j)
        cell.font = Font(size=10)
        cell.border = BORDE
        cell.alignment = Alignment(vertical='center', wrap_text=key in ('asegurado', 'situacion'))
        if key == 'saldo':
            cell.number_format = '"$" #,##0.00'
            cell.alignment = Alignment(horizontal='right', vertical='center')
        elif key == 'vencimiento' and cell.value:
            cell.number_format = 'dd/mm/yyyy'
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif key in ('dias', 'cuota', 'estado', 'productor'):
            cell.alignment = Alignment(horizontal='center', vertical='center')
        if key == 'wa_link' and cell.value:
            cell.hyperlink = cell.value
            cell.value = 'Abrir chat'
            cell.font = Font(size=10, color='0563C1', underline='single')
    if rec['estado'] == 'ANULADA':
        for j in range(1, len(COLS) + 1):
            ws2.cell(row=i, column=j).fill = PatternFill('solid', fgColor=GRIS)
            ws2.cell(row=i, column=j).font = Font(size=10, color='808080', italic=True)
    elif i % 2 == 0:
        for j in range(1, len(COLS) + 1):
            ws2.cell(row=i, column=j).fill = PatternFill('solid', fgColor='F7F9FC')

ws2.conditional_formatting.add(f'G2:G{last}',
    CellIsRule(operator='greaterThan', formula=['30'], fill=PatternFill('solid', fgColor=ROJO), font=Font(color='9C0006', bold=True)))
ws2.conditional_formatting.add(f'G2:G{last}',
    CellIsRule(operator='between', formula=['1', '30'], fill=PatternFill('solid', fgColor=AMBAR)))
ws2.conditional_formatting.add(f'G2:G{last}',
    CellIsRule(operator='lessThan', formula=['1'], fill=PatternFill('solid', fgColor=VERDE)))

tot = last + 1
ws2.cell(row=tot, column=1, value=f'TOTAL — {len(filas)} cuotas').font = Font(bold=True, size=11)
ws2.merge_cells(start_row=tot, start_column=1, end_row=tot, end_column=7)
ws2.cell(row=tot, column=1).alignment = Alignment(horizontal='right')
tc = ws2.cell(row=tot, column=8, value=total_global)
tc.number_format = '"$" #,##0.00'
tc.font = Font(bold=True, size=11, color=AZUL)
tc.alignment = Alignment(horizontal='right')
for j in range(1, len(COLS) + 1):
    ws2.cell(row=tot, column=j).fill = PatternFill('solid', fgColor=AZUL_CLARO)
    ws2.cell(row=tot, column=j).border = Border(top=Side(style='double', color=AZUL))

ws2.freeze_panes = 'A2'
ws2.auto_filter.ref = f'A1:{get_column_letter(len(COLS))}{last}'
ws2.sheet_view.zoomScale = 90
ws2.page_setup.orientation = 'landscape'
ws2.page_setup.fitToWidth = 1
ws2.sheet_properties.pageSetUpPr.fitToPage = True
ws2.print_title_rows = '1:1'

# =================== POR CLIENTE ===================
ws3 = wb.create_sheet('Por cliente')
ws3.sheet_properties.tabColor = '2E75B6'
cli = defaultdict(lambda: {'polizas': set(), 'cuotas': 0, 'total': 0.0, 'venc': None,
                           'tel': '', 'wa': '', 'email': '', 'mat': '', 'prod': '', 'nombre': ''})
for rec in filas:
    k = (rec['productor'], rec['matricula'])
    c = cli[k]
    c['prod'] = rec['productor']
    c['nombre'] = rec['asegurado']
    c['mat'] = rec['matricula']
    if rec['poliza']:
        c['polizas'].add(rec['poliza'])
    c['cuotas'] += 1
    c['total'] += rec['saldo']
    if rec['vencimiento'] and (c['venc'] is None or rec['vencimiento'] < c['venc']):
        c['venc'] = rec['vencimiento']
    c['tel'] = c['tel'] or rec['telefono']
    c['wa'] = c['wa'] or rec['whatsapp']
    c['email'] = c['email'] or rec['email']

cols3 = [('Productor', 11), ('Asegurado', 32), ('Matrícula', 11), ('Pólizas', 26), ('Cuotas', 9),
         ('Total adeudado', 15), ('Venc. más antiguo', 14), ('Días', 8),
         ('Teléfono', 15), ('WhatsApp', 16), ('Email', 30), ('Link WhatsApp', 16)]
ws3.append([c[0] for c in cols3])
for i, c in enumerate(cols3, start=1):
    ws3.column_dimensions[get_column_letter(i)].width = c[1]

orden3 = sorted(cli.values(), key=lambda c: -c['total'])
for c in orden3:
    dias = (HOY - c['venc']).days if c['venc'] else None
    ws3.append([c['prod'], c['nombre'], f"M-{c['mat']}", ', '.join(sorted(c['polizas'])), c['cuotas'],
                c['total'], c['venc'], dias, c['tel'], c['wa'], c['email'],
                f"https://wa.me/{c['wa'].replace('+','')}" if c['wa'].startswith('+549') else ''])

last3 = len(orden3) + 1
for c in ws3[1]:
    c.fill = PatternFill('solid', fgColor=AZUL)
    c.font = Font(bold=True, color='FFFFFF', size=10)
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws3.row_dimensions[1].height = 30
for i in range(2, last3 + 1):
    for j in range(1, len(cols3) + 1):
        cell = ws3.cell(row=i, column=j)
        cell.font = Font(size=10)
        cell.border = BORDE
        cell.alignment = Alignment(vertical='center', wrap_text=j in (2, 4))
        if j == 6:
            cell.number_format = '"$" #,##0.00'
            cell.alignment = Alignment(horizontal='right', vertical='center')
        elif j == 7:
            cell.number_format = 'dd/mm/yyyy'
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif j in (1, 3, 5, 8, 9):
            cell.alignment = Alignment(horizontal='center', vertical='center')
    if i % 2 == 0:
        for j in range(1, len(cols3) + 1):
            ws3.cell(row=i, column=j).fill = PatternFill('solid', fgColor='F7F9FC')
    wa_cell = ws3.cell(row=i, column=12)
    if wa_cell.value:
        wa_cell.hyperlink = wa_cell.value
        wa_cell.value = 'Abrir chat'
        wa_cell.font = Font(size=10, color='0563C1', underline='single')

ws3.conditional_formatting.add(f'H2:H{last3}',
    CellIsRule(operator='greaterThan', formula=['30'], fill=PatternFill('solid', fgColor=ROJO), font=Font(color='9C0006', bold=True)))
ws3.conditional_formatting.add(f'H2:H{last3}',
    CellIsRule(operator='lessThan', formula=['1'], fill=PatternFill('solid', fgColor=VERDE)))

tot3 = last3 + 1
ws3.cell(row=tot3, column=1, value=f'TOTAL — {len(orden3)} clientes').font = Font(bold=True, size=11)
ws3.merge_cells(start_row=tot3, start_column=1, end_row=tot3, end_column=5)
ws3.cell(row=tot3, column=1).alignment = Alignment(horizontal='right')
t3 = ws3.cell(row=tot3, column=6, value=sum(c['total'] for c in orden3))
t3.number_format = '"$" #,##0.00'
t3.font = Font(bold=True, size=11, color=AZUL)
t3.alignment = Alignment(horizontal='right')
for j in range(1, len(cols3) + 1):
    ws3.cell(row=tot3, column=j).fill = PatternFill('solid', fgColor=AZUL_CLARO)
    ws3.cell(row=tot3, column=j).border = Border(top=Side(style='double', color=AZUL))

ws3.freeze_panes = 'A2'
ws3.auto_filter.ref = f'A1:L{last3}'
ws3.sheet_view.zoomScale = 90
ws3.page_setup.orientation = 'landscape'
ws3.page_setup.fitToWidth = 1
ws3.sheet_properties.pageSetUpPr.fitToPage = True

os.makedirs(os.path.dirname(args.out) or '.', exist_ok=True)
wb.save(args.out)
print(f'OK -> {args.out} | {len(orden)} productores | {len(filas)} cuotas | {len(orden3)} clientes | total $ {ar(total_global)}')
for user, p in orden:
    print(f'   {user} {p["nombre"]}: {p["cuotas"]} cuotas | $ {ar(p["total"])}')
