# -*- coding: utf-8 -*-
import openpyxl

wb = openpyxl.load_workbook('deudas.xlsx')
ws = wb['Deudas']
h = [c.value for c in ws[1]]
i = {x: j for j, x in enumerate(h)}

def col(prefix):
    for name, idx in i.items():
        if name and name.lower().startswith(prefix.lower()):
            return idx
    raise KeyError(prefix)

SALDO, DIAS, ESTADO, WA, OTROS = col('Saldo'), col('D'), col('Estado'), col('Link'), col('Otros')

rows = list(ws.iter_rows(min_row=2, max_row=ws.max_row - 1, values_only=True))
tot = sum(r[SALDO] for r in rows)
venc = [r for r in rows if (r[DIAS] or 0) > 0]
tot_venc = sum(r[SALDO] for r in venc)
est = {}
for r in rows:
    est[r[ESTADO]] = est.get(r[ESTADO], 0) + 1

print('cuotas:', len(rows), '| clientes:', wb['Por cliente'].max_row - 2)
print('total: $ {:,.2f}'.format(tot))
print('vencidas: {} -> $ {:,.2f}'.format(len(venc), tot_venc))
print('a vencer: {} -> $ {:,.2f}'.format(len(rows) - len(venc), tot - tot_venc))
print('sin link WhatsApp:', sum(1 for r in rows if not r[WA]))
print('con otros emails:', sum(1 for r in rows if r[OTROS]))
print('estados:', est)
