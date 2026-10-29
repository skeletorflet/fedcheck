# -*- coding: utf-8 -*-
import openpyxl

for u in ['26620', '29494']:
    p = f'salida/{u}/deudas.xlsx'
    wb = openpyxl.load_workbook(p)
    print('==', p, '|', wb.sheetnames)
    ws = wb['Resumen']
    for row in ws.iter_rows(min_row=4, max_row=12, values_only=True):
        if any(x is not None for x in row):
            print('   ', [x for x in row if x is not None])
    d = wb['Deudas']
    print('    filas:', d.max_row - 1, '| total:', d.cell(row=d.max_row, column=7).value)
