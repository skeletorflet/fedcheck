# -*- coding: utf-8 -*-
import openpyxl

def resumen(path):
    wb = openpyxl.load_workbook(path)
    print('\n==', path)
    for ws in wb:
        print(f'  {ws.title}: {ws.max_row} filas x {ws.max_column} cols | filtro {ws.auto_filter.ref} | freeze {ws.freeze_panes}')

resumen('deudas_todas.xlsx')
wb = openpyxl.load_workbook('deudas_todas.xlsx')
ws = wb['Deudas']
h = [c.value for c in ws[1]]
print('  columnas:', h)
rows = list(ws.iter_rows(min_row=2, max_row=ws.max_row - 1, values_only=True))
print('  cuotas:', len(rows), '| total:', round(sum(r[7] for r in rows), 2))
print('  productores:', sorted({r[0] for r in rows}))
ws2 = wb['Por cliente']
rows2 = list(ws2.iter_rows(min_row=2, max_row=ws2.max_row - 1, values_only=True))
print('  clientes:', len(rows2), '| total:', round(sum(r[5] for r in rows2), 2))

ws3 = wb['Resumen']
print('  resumen:')
for row in ws3.iter_rows(min_row=4, max_row=9, values_only=True):
    print('   ', [x for x in row if x is not None])

# formulas?
formulas = 0
for w in wb:
    for row in w.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('='):
                formulas += 1
print('  formulas:', formulas)
